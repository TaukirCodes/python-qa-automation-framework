from pathlib import Path

import pytest
from openpyxl import load_workbook

from pages.login_page import LoginPage


DATA_FILE = (
    Path(__file__).resolve().parents[2]
    / "test_data"
    / "login_data.xlsx"
)


def load_excel_data():
    workbook = load_workbook(
        DATA_FILE,
        read_only=True,
        data_only=True
    )

    try:
        sheet = workbook["LoginData"]
        rows = sheet.iter_rows(values_only=True)

        headers = next(rows)
        test_cases = []

        for values in rows:
            if all(value is None for value in values):
                continue

            row = dict(zip(headers, values))

            test_cases.append(
                pytest.param(
                    row["username"],
                    row["password"],
                    row["expected_error"],
                    id=str(row["test_id"])
                )
            )

        return test_cases

    finally:
        workbook.close()


@pytest.mark.parametrize(
    "username,password,expected_error",
    load_excel_data()
)
def test_invalid_login_excel(
    driver, username, password, expected_error
):
    page = LoginPage(driver)

    page.open()
    page.login(username, password)

    actual_error = page.get_error_message()

    assert actual_error == expected_error
